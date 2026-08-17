import os
import json
import uuid
import time
import psutil
import tempfile
import openpyxl
import threading
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from flask import Flask, request, jsonify, send_file, make_response
from flask_cors import CORS

app = Flask(__name__)

# Global lock for thread-safe JSON files reading/writing
FILE_LOCK = threading.Lock()

# Helper to locate the WAL file for a given try_id
def get_wal_path(try_id=None):
    """Return the path to the WAL file.
    If try_id is provided, use the per-try WAL inside its directory.
    Otherwise, use the legacy global WAL_FILE.
    """
    if try_id:
        return os.path.join(BASE_DIR, f"try_{try_id}", "wal.json")
    return WAL_FILE
CORS(app)  # Allow extension to hit the API

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "config.json")
WAL_FILE = os.path.join(BASE_DIR, "wal.json")
LOG_FILE = os.path.join(BASE_DIR, "events.log")

# In-memory cache for consolidated dataset
DATASET_CACHE = {}

# In-memory cache for loaded JSON files to prevent heavy disk read/write cycles
JSON_CACHE = {}

def log_event(event_type, details):
    timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
    log_entry = f"[{timestamp}] [{event_type.upper()}] {json.dumps(details)}\n"
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(log_entry)

def load_json(filepath, default):
    if not os.path.exists(filepath):
        return default
    try:
        mtime = os.path.getmtime(filepath)
    except Exception:
        mtime = 0

    if filepath in JSON_CACHE:
        cached_mtime, cached_data = JSON_CACHE[filepath]
        if cached_mtime == mtime:
            return cached_data

    for attempt in range(5):
        try:
            with FILE_LOCK:
                with open(filepath, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    JSON_CACHE[filepath] = (mtime, data)
                    return data
        except (PermissionError, json.JSONDecodeError):
            time.sleep(0.05 * (attempt + 1))
            continue
        except Exception:
            return default
    return default

def save_json(filepath, data):
    dir_name = os.path.dirname(filepath)
    if dir_name:
        os.makedirs(dir_name, exist_ok=True)
    fd, temp_path = tempfile.mkstemp(dir=dir_name or None, suffix=".tmp")
    try:
        with FILE_LOCK:
            with os.fdopen(fd, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2)
            os.replace(temp_path, filepath)
            try:
                mtime = os.path.getmtime(filepath)
                JSON_CACHE[filepath] = (mtime, data)
            except Exception:
                JSON_CACHE.pop(filepath, None)
    except Exception as e:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass
        raise e

def get_effective_config():
    defaults = {"cpu_tabs": 1, "ram_batch_size": 1, "cpu_threshold": 85, "max_retries": 3}
    config = load_json(CONFIG_FILE, defaults)
    for k, v in defaults.items():
        if k not in config:
            config[k] = v
    # Force batch size to 1 to ensure running tasks match tab capacity (1 task per tab)
    config["ram_batch_size"] = 1
    return config

@app.route('/config', methods=['GET'])
def get_config():
    return jsonify(get_effective_config())

@app.route('/config', methods=['POST'])
def set_config():
    data = request.json
    config = get_effective_config()
    if "cpu_tabs" in data: config["cpu_tabs"] = data["cpu_tabs"]
    if "ram_batch_size" in data: config["ram_batch_size"] = data["ram_batch_size"]
    if "cpu_threshold" in data: config["cpu_threshold"] = data["cpu_threshold"]
    if "max_retries" in data: config["max_retries"] = data["max_retries"]
    save_json(CONFIG_FILE, config)
    return jsonify({"success": True, "config": config})

@app.route('/start_try', methods=['POST'])
def start_try():
    data = request.json
    permutations = data.get('permutations', [])
    target_url = data.get('target_url', 'https://one.dat.com/search-loads')
    
    if not permutations:
        return jsonify({"success": False, "error": "No permutations provided"})

    config = get_effective_config()
    # Force batch size to 1 to process tasks strictly according to the tab capacity
    batch_size = 1

    # Chunk permutations into batches
    batches = []
    for i in range(0, len(permutations), batch_size):
        chunk = permutations[i:i + batch_size]
        batches.append({
            "id": f"batch_{i//batch_size + 1}",
            "tasks": chunk,
            "status": "pending",
            "retries": 0
        })

    try_id = str(uuid.uuid4())
    try_dir = os.path.join(BASE_DIR, f"try_{try_id}")
    os.makedirs(try_dir, exist_ok=True)

    wal = {
        "try_id": try_id,
        "target_url": target_url,
        "status": "in_progress",
        "batches": batches,
        "created_at": time.time()
    }
    # Save WAL both globally (for backward compatibility) and in the try directory
    save_json(WAL_FILE, wal)
    save_json(os.path.join(try_dir, "wal.json"), wal)
    
    # Invalidate dataset cache
    global DATASET_CACHE
    DATASET_CACHE.pop(try_id, None)

    log_event("START_TRY", {"try_id": try_id, "target_url": target_url, "total_batches": len(batches), "total_permutations": len(permutations)})

    return jsonify({"success": True, "try_id": try_id, "total_batches": len(batches)})

@app.route('/next_batch', methods=['GET'])
def next_batch():
    # Determine try_id if present in request args for robustness
    # Determine try_id if present in request args; if none, discover most recent try directory
    try_id = request.args.get('try_id') or (request.json.get('try_id') if request.is_json else None)
    if not try_id:
        # Find existing try directories
        try_dirs = [d for d in os.listdir(BASE_DIR) if d.startswith('try_') and os.path.isdir(os.path.join(BASE_DIR, d))]
        if try_dirs:
            # Choose most recent by wal.json modification time
            try_dirs_paths = [os.path.join(BASE_DIR, d) for d in try_dirs]
            latest_dir = max(try_dirs_paths, key=lambda p: os.path.getmtime(os.path.join(p, 'wal.json')) if os.path.exists(os.path.join(p, 'wal.json')) else 0)
            wal_path = os.path.join(latest_dir, 'wal.json')
            wal = load_json(wal_path, {})
            # Sync global WAL
            if wal:
                save_json(WAL_FILE, wal)
                try_id = wal.get('try_id')
        if not try_id:
            # If still no wal or try_id, fall back to global WAL (may be empty)
            try_id = None
    wal_path = get_wal_path(try_id)
    wal = load_json(wal_path, {})
    if not wal or wal.get("status") != "in_progress":
        return jsonify({"success": False, "error": "No active try"})

    # Find the next pending batch.
    for batch in wal.get("batches", []):
        if batch["status"] == "pending":
            batch["status"] = "in_progress"
            save_json(wal_path, wal)
            # Also sync global WAL for backward compatibility
            if wal_path != WAL_FILE:
                save_json(WAL_FILE, wal)
            return jsonify({
                "success": True, 
                "try_id": wal["try_id"], 
                "target_url": wal.get("target_url", "https://one.dat.com/search-loads"),
                "batch": batch
            })

    return jsonify({"success": False, "message": "No more pending batches"})

def normalize_load_record(r):
    # Check inside directoryData nested dictionary
    dir_data = r.get("directoryData")
    if isinstance(dir_data, dict):
        dir_data_mappings = {
            "docket": "dir_docket",
            "dot_number": "dir_dot_number",
            "entity_type": "dir_gen_entity_type",
            "coverage_to": "dir_active_ins_coverage_to",
            "credit_score": "dir_credit_score",
            "insurance_carrier": "dir_active_ins_carrier"
        }
        for k, target in dir_data_mappings.items():
            if not r.get(target) and dir_data.get(k):
                r[target] = dir_data[k]

    mappings = {
        "dir_docket": ["docket", "Docket"],
        "dir_dot_number": ["dot_number", "DotNumber"],
        "dir_gen_entity_type": ["entity_type", "EntityType"],
        "dir_active_ins_coverage_to": ["coverage_to", "CoverageTo"],
        "dir_credit_score": ["credit_score", "CreditScore"],
        "dir_active_ins_carrier": ["insurance_carrier", "InsuranceCarrier"]
    }
    for dir_key, fallback_keys in mappings.items():
        if not r.get(dir_key):
            for fk in fallback_keys:
                if r.get(fk):
                    r[dir_key] = r[fk]
                    break
    return r

@app.route('/complete_batch', methods=['POST'])
def complete_batch():
    data = request.json
    try_id = data.get("try_id")
    batch_id = data.get("batch_id")
    results = data.get("results", [])
    task_stats = data.get("task_stats", [])  # [{origin,dest,eq,lt,records,status}]

    wal_path = get_wal_path(try_id)
    wal = load_json(wal_path, {})
    if wal.get("try_id") != try_id:
        return jsonify({"success": False, "error": "Try ID mismatch"})

    # Normalize results before saving
    normalized_results = []
    for r in results:
        if isinstance(r, dict):
            normalized_results.append(normalize_load_record(dict(r)))
        else:
            normalized_results.append(r)

    # Save results
    try_dir = os.path.join(BASE_DIR, f"try_{try_id}")
    os.makedirs(try_dir, exist_ok=True)
    results_file = os.path.join(try_dir, f"{batch_id}.json")
    save_json(results_file, normalized_results)
    # (Skip redundant intermediate WAL write; written at end of complete_batch)


    # Save per-task stats file for visibility
    if task_stats:
        stats_file = os.path.join(try_dir, f"{batch_id}_stats.json")
        save_json(stats_file, task_stats)

    # Log skipped combinations separately (zero results, timeouts, equipment failures, search failures)
    skipped_tasks = [t for t in task_stats if t.get("status", "").startswith("skipped")]
    if skipped_tasks:
        skipped_file = os.path.join(try_dir, "skipped_combinations.json")
        existing_skipped = load_json(skipped_file, [])
        # Dedup by combination key
        existing_keys = set(f"{s.get('origin')}|{s.get('dest')}|{s.get('eq')}|{s.get('lt')}" for s in existing_skipped)
        for s in skipped_tasks:
            key = f"{s.get('origin')}|{s.get('dest')}|{s.get('eq')}|{s.get('lt')}"
            if key not in existing_keys:
                existing_skipped.append(s)
                existing_keys.add(key)
        save_json(skipped_file, existing_skipped)
        log_event("SKIPPED_COMBINATIONS_LOGGED", {"count": len(skipped_tasks), "try_id": try_id, "batch_id": batch_id})

    # Update WAL batch status and store task stats inline
    unfinished_tasks = []
    for batch in wal.get("batches", []):
        if batch["id"] == batch_id:
            completed_keys = {
                f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}"
                for ts in task_stats
            }
            for task in batch.get("tasks", []):
                eqs = task.get("equipment")
                if isinstance(eqs, list):
                    missing_eqs = []
                    for eq in eqs:
                        key = f"{task.get('origin')}|{task.get('dest')}|{eq}|{task.get('lt')}"
                        if key not in completed_keys:
                            missing_eqs.append(eq)
                    if missing_eqs:
                        rescheduled_task = dict(task)
                        rescheduled_task["equipment"] = missing_eqs
                        unfinished_tasks.append(rescheduled_task)
                else:
                    eq_val = task.get("eq") or task.get("equipment") or ""
                    key = f"{task.get('origin')}|{task.get('dest')}|{eq_val}|{task.get('lt')}"
                    if key not in completed_keys:
                        unfinished_tasks.append(task)
            
            batch["status"] = "completed"
            batch["records_scraped"] = len(results)
            batch["task_stats"] = task_stats
            batch["completed_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
            break

    # If there were unfinished tasks (e.g. user clicked Stop mid-batch), reschedule them in a new pending batch
    if unfinished_tasks:
        new_batch_id = f"batch_{len(wal.get('batches', [])) + 1}"
        wal.setdefault("batches", []).append({
            "id": new_batch_id,
            "tasks": unfinished_tasks,
            "status": "pending",
            "retries": 0
        })
        log_event("PARTIAL_BATCH_RESCHEDULED", {
            "original_batch": batch_id,
            "new_batch": new_batch_id,
            "rescheduled_count": len(unfinished_tasks)
        })

    # Check if all completed (either "completed" or "failed")
    all_finished = all(b.get("status") in ["completed", "failed"] for b in wal.get("batches", []))
    if all_finished:
        wal["status"] = "completed"

    # Write checkpoint
    wal["last_checkpoint"] = {
        "batch_id": batch_id,
        "records_scraped": len(results),
        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S")
    }

    save_json(wal_path, wal)
    if wal_path != WAL_FILE:
        save_json(WAL_FILE, wal)
    
    # Incrementally update dataset cache if it exists
    global DATASET_CACHE
    if try_id in DATASET_CACHE:
        cache_data = DATASET_CACHE[try_id]["data"]
        
        # 1. Update records
        cache_data["records"].extend(normalized_results)
        cache_data["counts"]["success_records"] = len(cache_data["records"])
        
        # 2. Update task_stats
        current_stats = cache_data.setdefault("task_stats", [])
        completed_keys = {
            f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}"
            for ts in task_stats
        }
        new_stats = [ts for ts in current_stats if f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}" not in completed_keys]
        new_stats.extend(task_stats)
        cache_data["task_stats"] = new_stats
        
        # 3. Update counts
        success_count = sum(1 for ts in new_stats if ts.get("status") == "completed")
        skipped_count = sum(1 for ts in new_stats if "skipped" in ts.get("status", ""))
        failed_count = sum(1 for ts in new_stats if ts.get("status") in ["error", "failed"])
        
        cache_data["counts"]["success_combinations"] = success_count
        cache_data["counts"]["skipped_combinations"] = skipped_count
        cache_data["counts"]["failed_combinations"] = failed_count
        
        DATASET_CACHE[try_id]["timestamp"] = time.time()

    log_event("BATCH_COMPLETED", {"try_id": try_id, "batch_id": batch_id, "extracted": len(results), "task_stats": task_stats})
    if all_finished:
        log_event("TRY_COMPLETED", {"try_id": try_id})

    return jsonify({"success": True, "all_completed": all_finished})

@app.route('/status', methods=['GET'])
def status():
    # Try to obtain try_id from query for proper WAL location
    try_id = request.args.get('try_id')
    wal_path = get_wal_path(try_id)
    wal = load_json(wal_path, {})
    if not wal:
        return jsonify({"active": False})

    batches = wal.get("batches", [])
    total = len(batches)
    completed = sum(1 for b in batches if b["status"] == "completed")
    pending = sum(1 for b in batches if b["status"] == "pending")
    in_progress = sum(1 for b in batches if b["status"] == "in_progress")
    failed = sum(1 for b in batches if b["status"] == "failed")

    # Monitor CPU and memory usage
    try:
        cpu_usage = psutil.cpu_percent(interval=None)
        memory_usage = psutil.virtual_memory().percent
    except Exception:
        cpu_usage = 0
        memory_usage = 0

    # Calculate task stats
    total_tasks = 0
    completed_tasks = 0
    skipped_tasks = 0
    failed_tasks = 0
    running_tasks = 0
    success_records = 0

    for b in batches:
        batch_tasks_count = 0
        for task in b.get("tasks", []):
            eqs = task.get("equipment")
            if isinstance(eqs, list):
                batch_tasks_count += len(eqs)
            else:
                batch_tasks_count += 1
                
        total_tasks += batch_tasks_count
        b_status = b.get("status")
        
        if b_status == "completed":
            success_records += b.get("records_scraped", 0)
            task_stats = b.get("task_stats", [])
            for ts in task_stats:
                status_val = ts.get("status")
                if status_val == "completed":
                    completed_tasks += 1
                elif "skipped" in status_val or status_val == "skipped_zero_results":
                    skipped_tasks += 1
                elif status_val == "error" or status_val == "failed":
                    failed_tasks += 1
        elif b_status == "failed":
            failed_tasks += batch_tasks_count
        elif b_status == "in_progress":
            running_tasks += batch_tasks_count
            
    remaining_tasks = total_tasks - completed_tasks - skipped_tasks - failed_tasks - running_tasks

    # Calculate ETA
    eta_str = "Calculating..."
    if completed > 0 and total > 0:
        elapsed = time.time() - wal.get("created_at", time.time())
        avg_time = elapsed / completed
        remaining_batches = total - completed
        eta_sec = avg_time * remaining_batches
        if eta_sec < 60:
            eta_str = f"{int(eta_sec)}s"
        elif eta_sec < 3600:
            eta_str = f"{int(eta_sec // 60)}m {int(eta_sec % 60)}s"
        else:
            eta_str = f"{int(eta_sec // 3600)}h {int((eta_sec % 3600) // 60)}m"
    elif total == completed and total > 0:
        eta_str = "Completed"

    config = get_effective_config()
    concurrency_limit = config.get("cpu_tabs", 1)

    return jsonify({
        "active": True,
        "try_id": wal.get("try_id"),
        "status": wal.get("status"),
        "total": total,
        "completed": completed,
        "pending": pending,
        "in_progress": in_progress if wal.get("status") == "in_progress" else 0,
        "failed": failed,
        "cpu_usage": cpu_usage,
        "memory_usage": memory_usage,
        "eta": eta_str,
        "concurrency_limit": concurrency_limit,
        "tasks": {
            "total": total_tasks,
            "completed": completed_tasks,
            "skipped": skipped_tasks,
            "failed": failed_tasks,
            "remaining": remaining_tasks,
            "running": running_tasks if wal.get("status") == "in_progress" else 0,
            "success_records": success_records
        }
    })

@app.route('/logs', methods=['GET'])
def get_logs():
    if not os.path.exists(LOG_FILE):
        return jsonify([])
    try:
        with open(LOG_FILE, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        parsed_logs = []
        for line in lines[-150:]:  # last 150 events
            parts = line.strip().split(' ', 2)
            if len(parts) >= 3:
                timestamp = (parts[0] + ' ' + parts[1]).strip('[]')
                rest = parts[2].split('] ', 1)
                if len(rest) == 2:
                    event_type = rest[0].strip('[]')
                    try:
                        details = json.loads(rest[1])
                    except Exception:
                        details = rest[1]
                    parsed_logs.append({
                        "timestamp": timestamp,
                        "event_type": event_type,
                        "details": details
                    })
        return jsonify(parsed_logs)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/log_event', methods=['POST'])
def log_event_endpoint():
    data = request.json
    event_type = data.get("event_type", "UNKNOWN")
    details = data.get("details", {})
    log_event(event_type, details)
    return jsonify({"success": True})

@app.route('/reset_batch', methods=['POST'])
def reset_batch():
    """Reset an in_progress batch back to pending (for crash recovery or timeout)."""
    data = request.get_json(silent=True) or {}
    batch_id = data.get("batch_id")
    try_id = data.get("try_id")
    if not try_id:
        try_dirs = [d for d in os.listdir(BASE_DIR) if d.startswith('try_') and os.path.isdir(os.path.join(BASE_DIR, d))]
        if try_dirs:
            try_dirs_paths = [os.path.join(BASE_DIR, d) for d in try_dirs]
            latest_dir = max(try_dirs_paths, key=lambda p: os.path.getmtime(os.path.join(p, 'wal.json')) if os.path.exists(os.path.join(p, 'wal.json')) else 0)
            wal_path = os.path.join(latest_dir, 'wal.json')
            wal = load_json(wal_path, {})
            if wal:
                try_id = wal.get('try_id')
    reason = data.get("reason", "tab_closed_unexpectedly")
    
    config = get_effective_config()
    max_retries = config.get("max_retries", 3)
    
    wal_path = get_wal_path(try_id)
    wal = load_json(wal_path, {})
    
    for batch in wal.get("batches", []):
        if batch["id"] == batch_id:
            retries = batch.get("retries", 0)
            if retries >= max_retries:
                batch["status"] = "failed"
                batch["error"] = f"Max retries ({max_retries}) exceeded"
                save_json(wal_path, wal)
                if wal_path != WAL_FILE: save_json(WAL_FILE, wal)
                if try_id:
                    DATASET_CACHE.pop(try_id, None)
                log_event("BATCH_FAILED", {"batch_id": batch_id, "retries": retries, "reason": "max_retries_exceeded"})
                return jsonify({"success": False, "error": "Max retries exceeded. Batch marked as failed."})
            else:
                batch["status"] = "pending"
                batch["retries"] = retries + 1
                save_json(wal_path, wal)
                if wal_path != WAL_FILE: save_json(WAL_FILE, wal)
                if try_id:
                    DATASET_CACHE.pop(try_id, None)
                log_event("BATCH_RESET", {"batch_id": batch_id, "attempt": retries + 1, "reason": reason})
                return jsonify({"success": True, "attempt": retries + 1})
                
    return jsonify({"success": False, "error": "Batch not found"})

@app.route('/stop_try', methods=['POST'])
def stop_try():
    data = request.get_json(silent=True) or {}
    try_id = data.get("try_id")
    
    wal = None
    wal_path = None
    if try_id and try_id != "Starting...":
        wal_path = get_wal_path(try_id)
        if os.path.exists(wal_path):
            wal = load_json(wal_path, {})
            
    if not wal:
        # Fall back to finding the latest try directory
        try_dirs = [d for d in os.listdir(BASE_DIR) if d.startswith('try_') and os.path.isdir(os.path.join(BASE_DIR, d))]
        if try_dirs:
            try_dirs_paths = [os.path.join(BASE_DIR, d) for d in try_dirs]
            latest_dir = max(try_dirs_paths, key=lambda p: os.path.getmtime(os.path.join(p, 'wal.json')) if os.path.exists(os.path.join(p, 'wal.json')) else 0)
            wal_path = os.path.join(latest_dir, 'wal.json')
            wal = load_json(wal_path, {})
            if wal:
                try_id = wal.get('try_id')
                
    if not wal:
        # Fall back to root WAL file
        wal_path = WAL_FILE
        if os.path.exists(wal_path):
            wal = load_json(wal_path, {})
            
    if wal:
        wal["status"] = "cancelled"
        
        # Reset any batches left in-progress back to pending
        for batch in wal.get("batches", []):
            if batch.get("status") == "in_progress":
                batch["status"] = "pending"
                
        save_json(wal_path, wal)
        if wal_path != WAL_FILE: save_json(WAL_FILE, wal)
        log_event("TRY_CANCELLED", {"try_id": wal.get("try_id")})
        
        # Invalidate dataset cache
        actual_try_id = wal.get("try_id")
        global DATASET_CACHE
        DATASET_CACHE.pop(actual_try_id, None)
        
        return jsonify({"success": True})
    return jsonify({"success": True, "message": "No active try found"})



@app.route('/resume', methods=['POST'])
def resume():
    """Resume the most recent extraction session, resetting in-progress batches."""
    # Discover existing try directories
    try_dirs = [d for d in os.listdir(BASE_DIR) if d.startswith('try_') and os.path.isdir(os.path.join(BASE_DIR, d))]
    if not try_dirs:
        return jsonify({"success": False, "error": "No previous tries found"})
    try_dirs_paths = [os.path.join(BASE_DIR, d) for d in try_dirs]
    latest_dir = max(try_dirs_paths, key=lambda p: os.path.getmtime(os.path.join(p, 'wal.json')) if os.path.exists(os.path.join(p, 'wal.json')) else 0)
    wal_path = os.path.join(latest_dir, 'wal.json')
    wal = load_json(wal_path, {})
    if not wal:
        return jsonify({"success": False, "error": "WAL not found for latest try"})
    try_id = wal.get('try_id')
    wal['status'] = 'in_progress'
    # Reset any in_progress batches to pending
    reset_count = 0
    for batch in wal.get('batches', []):
        if batch['status'] == 'in_progress':
            batch['status'] = 'pending'
            reset_count += 1
    save_json(wal_path, wal)
    # Sync global WAL
    if wal_path != WAL_FILE:
        save_json(WAL_FILE, wal)
    # Invalidate cache
    global DATASET_CACHE
    DATASET_CACHE.pop(try_id, None)
    log_event('TRY_RESUMED', {'try_id': try_id, 'reset_batches': reset_count})
    return jsonify({"success": True, "try_id": try_id, "reset_batches": reset_count})

@app.route('/dataset', methods=['GET'])
def get_dataset():
    wal = load_json(WAL_FILE, {})
    try_id = wal.get("try_id")
    if not try_id:
        return jsonify({"success": False, "error": "No active try found"})
        
    global DATASET_CACHE
    if try_id in DATASET_CACHE:
        return jsonify(DATASET_CACHE[try_id]["data"])
        
    try_dir = os.path.join(BASE_DIR, f"try_{try_id}")
    if not os.path.exists(try_dir):
        return jsonify({"success": True, "records": [], "task_stats": []})
        
    all_records = []
    task_stats = []
    
    for f in sorted(os.listdir(try_dir)):
        if f.endswith("_stats.json"):
            stats_data = load_json(os.path.join(try_dir, f), [])
            task_stats.extend(stats_data)
        elif f.endswith(".json") and not f.startswith("skipped_") and f != "dataset.json" and f != "dataset_stats.json" and f != "wal.json":
            batch_data = load_json(os.path.join(try_dir, f), [])
            all_records.extend(batch_data)
            
    skipped_file = os.path.join(try_dir, "skipped_combinations.json")
    if os.path.exists(skipped_file):
        skipped_data = load_json(skipped_file, [])
        seen = set()
        unique_stats = []
        for ts in task_stats:
            key = f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}"
            if key not in seen:
                seen.add(key)
                unique_stats.append(ts)
        for ts in skipped_data:
            key = f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}"
            if key not in seen:
                seen.add(key)
                unique_stats.append(ts)
        task_stats = unique_stats

    # Count combination statuses
    success_count = sum(1 for ts in task_stats if ts.get("status") == "completed")
    skipped_count = sum(1 for ts in task_stats if "skipped" in ts.get("status", ""))
    failed_count = sum(1 for ts in task_stats if ts.get("status") in ["error", "failed"])

    response_data = {
        "success": True,
        "try_id": try_id,
        "records": all_records,
        "task_stats": task_stats,
        "counts": {
            "success_records": len(all_records),
            "success_combinations": success_count,
            "skipped_combinations": skipped_count,
            "failed_combinations": failed_count
        }
    }
    
    DATASET_CACHE[try_id] = {
        "timestamp": time.time(),
        "data": response_data
    }
    return jsonify(response_data)

def get_export_data(try_id, export_type):
    try_dir = os.path.join(BASE_DIR, f"try_{try_id}")
    if not os.path.exists(try_dir):
        return []
        
    all_records = []
    task_stats = []
    
    for f in sorted(os.listdir(try_dir)):
        if f.endswith("_stats.json"):
            stats_data = load_json(os.path.join(try_dir, f), [])
            task_stats.extend(stats_data)
        elif f.endswith(".json") and not f.startswith("skipped_") and f != "dataset.json" and f != "dataset_stats.json" and f != "wal.json":
            batch_data = load_json(os.path.join(try_dir, f), [])
            all_records.extend(batch_data)
            
    skipped_file = os.path.join(try_dir, "skipped_combinations.json")
    if os.path.exists(skipped_file):
        skipped_data = load_json(skipped_file, [])
        seen = set()
        unique_stats = []
        for ts in task_stats:
            key = f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}"
            if key not in seen:
                seen.add(key)
                unique_stats.append(ts)
        for ts in skipped_data:
            key = f"{ts.get('origin')}|{ts.get('dest')}|{ts.get('eq')}|{ts.get('lt')}"
            if key not in seen:
                seen.add(key)
                unique_stats.append(ts)
        task_stats = unique_stats

    rows = []
    
    dir_fields = [
        "dir_docket", "dir_dot_number", "dir_gen_entity_type", "dir_active_ins_coverage_to", "dir_credit_score", "dir_active_ins_carrier"
    ]

    # 1. Success records
    if export_type in ["all", "success"]:
        for r in all_records:
            row_dict = {
                # 32 keys
                "postingId": r.get("postingId", ""),
                "timestamp": r.get("timestamp", time.strftime("%Y-%m-%d %H:%M:%S")),
                "age": r.get("age", "–"),
                "rate": r.get("rate", "–"),
                "ratePerMile": r.get("ratePerMile", "–"),
                "tripMiles": r.get("tripMiles", "–"),
                "origin": r.get("origin", "–"),
                "destination": r.get("destination", "–"),
                "deadheadOrigin": r.get("deadheadOrigin", ""),
                "deadheadDest": r.get("deadheadDest", ""),
                "equipmentType": r.get("equipmentType", "–"),
                "weight": r.get("weight", "–"),
                "length": r.get("length", "–"),
                "loadType": r.get("loadType", "–"),
                "company": r.get("company", "–"),
                "truckType": r.get("truckType", ""),
                "factoring": r.get("factoring", "No"),
                "referenceId": r.get("referenceId", ""),
                "pickupDate": r.get("pickupDate", ""),
                "pickupTime": r.get("pickupTime", ""),
                "deliveryTime": r.get("deliveryTime", ""),
                "phone": r.get("phone", ""),
                "email": r.get("email", ""),
                "companyLocation": r.get("companyLocation", ""),
                "isEmpty": r.get("isEmpty", False),
                "rowStatus": r.get("rowStatus", "Success"),
                "_captureKey": r.get("_captureKey", ""),
                "listIndex": r.get("listIndex", ""),
                "_apOrigin": r.get("_apOrigin", ""),
                "_apDest": r.get("_apDest", ""),
                "_apEq": r.get("_apEq", ""),
                "_apLt": r.get("_apLt", ""),
                
                # Old keys (for backward compatibility with CSV/TXT/Status)
                "rpm": r.get("ratePerMile", "–"),
                "trip": r.get("tripMiles", "–"),
                "equipment": r.get("equipmentType", "–"),
                "dho": r.get("deadheadOrigin", ""),
                "dhd": r.get("deadheadDest", ""),
                "load_type": r.get("loadType", "–"),
                "ref_id": r.get("referenceId", ""),
                "status": "Success",
                "error": ""
            }
            # Add directory fields
            for df in dir_fields:
                row_dict[df] = r.get(df, "")
            # Deep-extraction payloads ride along nested; exporters flatten on demand
            # and the JSON export keeps the full structure untouched.
            row_dict["_deep"] = r.get("deep")
            row_dict["_extraction"] = r.get("extraction")
            rows.append(row_dict)
            
    # 2. Skipped records
    if export_type in ["all", "skipped"]:
        for ts in task_stats:
            if "skipped" in ts.get("status", ""):
                status_label = "Skipped (0 results)" if ts.get("status") == "skipped_zero_results" else f"Skipped ({ts.get('status')})"
                row_dict = {
                    # 32 keys
                    "postingId": "",
                    "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                    "age": "–",
                    "rate": "–",
                    "ratePerMile": "–",
                    "tripMiles": "–",
                    "origin": ts.get("origin", "–"),
                    "destination": ts.get("dest", "–"),
                    "deadheadOrigin": "",
                    "deadheadDest": "",
                    "equipmentType": ts.get("eq", "–"),
                    "weight": "–",
                    "length": "–",
                    "loadType": ts.get("lt", "–"),
                    "company": "–",
                    "truckType": "",
                    "factoring": "No",
                    "referenceId": "",
                    "pickupDate": "",
                    "pickupTime": "",
                    "deliveryTime": "",
                    "phone": "",
                    "email": "",
                    "companyLocation": "",
                    "isEmpty": True,
                    "rowStatus": status_label,
                    "_captureKey": "",
                    "listIndex": "",
                    "_apOrigin": ts.get("origin", ""),
                    "_apDest": ts.get("dest", ""),
                    "_apEq": ts.get("eq", ""),
                    "_apLt": ts.get("lt", ""),
                    
                    # Old keys
                    "rpm": "–",
                    "trip": "–",
                    "equipment": ts.get("eq", "–"),
                    "dho": "",
                    "dhd": "",
                    "load_type": ts.get("lt", "–"),
                    "ref_id": "",
                    "status": status_label,
                    "error": ""
                }
                for df in dir_fields:
                    row_dict[df] = ""
                rows.append(row_dict)
                
    # 3. Failed records
    if export_type in ["all", "failed"]:
        for ts in task_stats:
            if ts.get("status") in ["error", "failed"]:
                row_dict = {
                    # 32 keys
                    "postingId": "",
                    "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                    "age": "–",
                    "rate": "–",
                    "ratePerMile": "–",
                    "tripMiles": "–",
                    "origin": ts.get("origin", "–"),
                    "destination": ts.get("dest", "–"),
                    "deadheadOrigin": "",
                    "deadheadDest": "",
                    "equipmentType": ts.get("eq", "–"),
                    "weight": "–",
                    "length": "–",
                    "loadType": ts.get("lt", "–"),
                    "company": "–",
                    "truckType": "",
                    "factoring": "No",
                    "referenceId": "",
                    "pickupDate": "",
                    "pickupTime": "",
                    "deliveryTime": "",
                    "phone": "",
                    "email": "",
                    "companyLocation": "",
                    "isEmpty": True,
                    "rowStatus": "Failed",
                    "_captureKey": "",
                    "listIndex": "",
                    "_apOrigin": ts.get("origin", ""),
                    "_apDest": ts.get("dest", ""),
                    "_apEq": ts.get("eq", ""),
                    "_apLt": ts.get("lt", ""),
                    
                    # Old keys
                    "rpm": "–",
                    "trip": "–",
                    "equipment": ts.get("eq", "–"),
                    "dho": "",
                    "dhd": "",
                    "load_type": ts.get("lt", "–"),
                    "ref_id": "",
                    "status": "Failed",
                    "error": ts.get("error", "Unknown error")
                }
                for df in dir_fields:
                    row_dict[df] = ""
                rows.append(row_dict)
                
    return rows

# ── Deep extraction export layer ─────────────────────────────────────────────
# Mirrors DEEP_COLUMNS in dat-load-extractor-v2/deep-extract.js. Keep the two
# lists in step: same headers, same order, so a CSV produced by the manual
# extractor and one produced by the distributed autopilot are interchangeable.
DEEP_EXPORT_COLUMNS = [
    ("Load ID", "details.loadId"),
    ("Load Number", "details.loadNumber"),
    ("Load Status", "details.loadStatus"),
    ("Detail Load Type", "details.loadType"),
    ("Commodity", "details.commodity"),
    ("Detail Reference ID", "details.referenceId"),
    ("Detail Pickup Date", "details.pickupDate"),
    ("Detail Delivery Date", "details.deliveryDate"),
    ("Detail Pickup Time", "details.pickupTime"),
    ("Detail Delivery Time", "details.deliveryTime"),
    ("Notes", "details.notes"),

    ("Route Origin", "route.origin"),
    ("Route Destination", "route.destination"),
    ("Route Text", "route.routeText"),
    ("Route Stops", "route.stopsFlat"),
    ("Route Stop Count", "route.stopCount"),
    ("Route Total Miles", "route.totalMiles"),
    ("Route Distance", "route.routeDistance"),
    ("Route DH Origin", "route.deadheadOrigin"),
    ("Route DH Destination", "route.deadheadDestination"),

    ("Truck", "truck.truckType"),
    ("Detail Equipment", "truck.equipmentType"),
    ("Trailer", "truck.trailerType"),
    ("Detail Length", "truck.length"),
    ("Detail Weight", "truck.weight"),
    ("Capacity", "truck.capacity"),
    ("Dimensions", "truck.dimensions"),
    ("Special Requirements", "truck.specialRequirements"),

    ("Rate Total", "rate.rateTotal"),
    ("Total Trip", "rate.totalTrip"),
    ("Detail Rate/Mile", "rate.ratePerMile"),
    ("Total Cost", "rate.totalCost"),
    ("Spot Rate", "rate.spotRate"),
    ("Spot Rate/Mile", "rate.spotRatePerMile"),
    ("Spot Rate Basis", "rate.spotRateBasis"),
    ("Rate Market", "rate.rateMarket"),
    ("Rate Range", "rate.rateRange"),
    ("Rate Range Low", "rate.rateRangeLow"),
    ("Rate Range High", "rate.rateRangeHigh"),
    ("Rate Range/Mile Low", "rate.rateRangePerMileLow"),
    ("Rate Range/Mile High", "rate.rateRangePerMileHigh"),
    ("Contract Rate", "rate.contractRate"),
    ("Contract Rate Note", "rate.contractRateNote"),
    ("Line Haul", "rate.lineHaul"),
    ("Fuel Cost", "rate.fuelCost"),
    ("Accessorials", "rate.accessorials"),
    ("Detention", "rate.detention"),
    ("Tolls", "rate.tolls"),
    ("Other Charges", "rate.otherCharges"),

    ("Detail Company", "company.name"),
    ("Company Name On Record", "company.nameOnRecord"),
    ("DBA Name", "company.dbaName"),
    ("Authority Name", "company.authorityName"),
    ("MC Number", "company.mcNumber"),
    ("Detail Docket", "company.docket"),
    ("Detail DOT Number", "company.dotNumber"),
    ("Detail Credit Score", "company.creditScore"),
    ("Days To Pay", "company.daysToPay"),
    ("Company Type", "company.companyType"),
    ("Entity Type", "company.entityType"),
    ("Operating Status", "company.operatingStatus"),
    ("Operation Type", "company.operationType"),
    ("Parent Account", "company.parentAccount"),
    ("Affiliations", "company.affiliations"),
    ("Insurance Carrier", "company.insuranceCarrier"),
    ("Insurance Coverage To", "company.insuranceCoverageTo"),
    ("Insurance Coverage From", "company.insuranceCoverageFrom"),
    ("Insurance Policy", "company.insurancePolicy"),
    ("Insurance Effective Date", "company.insuranceEffectiveDate"),
    ("Insurance Cancellation Date", "company.insuranceCancellationDate"),
    ("Insurance Address", "company.insuranceAddress"),
    ("Insurance Phone", "company.insurancePhone"),
    ("Insurance Fax", "company.insuranceFax"),
    ("Business Phone", "company.businessPhone"),
    ("Business Fax", "company.businessFax"),
    ("Company Location", "company.location"),
    ("Company Rating", "company.rating"),

    ("Contact Name", "contact.name"),
    ("Contact Title", "contact.title"),
    ("Contact Phone", "contact.phone"),
    ("Contact Mobile", "contact.mobile"),
    ("Contact Fax", "contact.fax"),
    ("Contact Email", "contact.email"),
    ("Contact Website", "contact.website"),
    ("All Emails", "contact.emailsFlat"),
    ("All Phones", "contact.phonesFlat"),

    ("Office Name", "office.name"),
    ("Office Address", "office.addressRaw"),
    ("Office Address Line1", "office.address.line1"),
    ("Office Address Line2", "office.address.line2"),
    ("Office City", "office.address.city"),
    ("Office State", "office.address.state"),
    ("Office ZIP", "office.address.zip"),
    ("Office Country", "office.address.country"),
    ("Office Phone", "office.phone"),
    ("Office Fax", "office.fax"),
    ("Office Email", "office.email"),
    ("Office Hours", "office.hours"),
    ("Business Address", "office.businessAddress"),
    ("Mailing Address", "office.mailingAddress"),

    ("Extraction Status", "_ex.status"),
    ("Extraction Fields", "_ex.fieldsExtracted"),
    ("Sections Visited", "_ex.sectionsVisitedFlat"),
    ("Sections Failed", "_ex.sectionsFailedFlat"),
    ("Extraction Retries", "_ex.retries"),
    ("Extraction Warnings", "_ex.warningsFlat"),
    ("Additional Data", "_ex.additionalDataFlat"),
]

DEEP_EXPORT_HEADERS = [h for h, _ in DEEP_EXPORT_COLUMNS]


def _dig(obj, path):
    node = obj
    for part in path.split("."):
        if not isinstance(node, dict):
            return None
        node = node.get(part)
    return node


def deep_value(row, path):
    """Resolve one flattened deep-extraction column for an export row."""
    deep = row.get("_deep") or {}
    ex = row.get("_extraction") or {}

    if path.startswith("_ex."):
        key = path[4:]
        if key == "sectionsVisitedFlat":
            return " | ".join(ex.get("sectionsVisited") or [])
        if key == "sectionsFailedFlat":
            return " | ".join(ex.get("sectionsFailed") or [])
        if key == "warningsFlat":
            return " | ".join((ex.get("warnings") or [])[:5])
        if key == "additionalDataFlat":
            extra = deep.get("additionalData") or {}
            return " | ".join(f"{k}={v}" for k, v in extra.items())
        val = ex.get(key)
        return "" if val is None else str(val)

    if not deep:
        return ""

    if path == "route.stopsFlat":
        stops = (deep.get("route") or {}).get("stops") or []
        parts = []
        for s in stops:
            dh = f" ({s['deadheadMiles']})" if s.get("deadheadMiles") is not None else ""
            date = f" {s['date']}" if s.get("date") else ""
            parts.append(f"{s.get('location', '')}{dh}{date}".strip())
        return " -> ".join(parts)
    if path == "route.stopCount":
        stops = (deep.get("route") or {}).get("stops") or []
        return str(len(stops)) if stops else ""
    if path == "contact.emailsFlat":
        return "; ".join((deep.get("contact") or {}).get("emails") or [])
    if path == "contact.phonesFlat":
        return "; ".join((deep.get("contact") or {}).get("phones") or [])

    val = _dig(deep, path)
    if val is None:
        return ""
    if isinstance(val, list):
        return "; ".join(str(v) for v in val)
    if isinstance(val, dict):
        return json.dumps(val)
    return str(val)


def rows_have_deep(rows):
    return any(r.get("_deep") or r.get("_extraction") for r in rows)


def get_export_csv(rows):
    import io
    import csv
    output = io.StringIO()
    writer = csv.writer(output, lineterminator='\n')
    
    headers = [
        "Origin", "Destination", "Age", "Rate", "Rate/Mile", "Trip", 
        "Company", "Equipment", "DH Origin", "DH Destination", 
        "Weight", "Length", "Load Type", "Reference ID", "Phone", 
        "Email", "Factoring", "Status", "Timestamp", "Error",
        "Dir Docket", "Dir DOT Number", "Dir Entity Type", "Dir Coverage To", "Dir Credit Score", "Dir Insurance Carrier"
    ]
    include_deep = rows_have_deep(rows)
    if include_deep:
        headers = headers + DEEP_EXPORT_HEADERS
    writer.writerow(headers)

    for r in rows:
        row = [
            r["origin"], r["destination"], r["age"], r["rate"], r["rpm"], r["trip"],
            r["company"], r["equipment"], r["dho"], r["dhd"], r["weight"], r["length"],
            r["load_type"], r["ref_id"], r["phone"], r["email"], r["factoring"],
            r["status"], r["timestamp"], r["error"],
            r.get("dir_docket", ""), r.get("dir_dot_number", ""), r.get("dir_gen_entity_type", ""), r.get("dir_active_ins_coverage_to", ""), r.get("dir_credit_score", ""), r.get("dir_active_ins_carrier", "")
        ]
        if include_deep:
            row.extend(deep_value(r, path) for _, path in DEEP_EXPORT_COLUMNS)
        writer.writerow(row)
    return output.getvalue()

def get_export_txt(rows, export_type):
    lines = []
    lines.append("="*80)
    lines.append(f"DAT ONE TMS — LOAD EXPORT ({export_type.upper()})")
    lines.append(f"Exported: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"Total Rows: {len(rows)}")
    lines.append("="*80)
    lines.append("")
    
    for idx, r in enumerate(rows):
        lines.append(f"┌─ ROW #{idx + 1} [{r['status'].upper()}] " + "─"*40)
        lines.append(f"│  Origin:        {r['origin']}")
        lines.append(f"│  Destination:   {r['destination']}")
        lines.append(f"│  Equipment:     {r['equipment']}")
        lines.append(f"│  Load Type:     {r['load_type']}")
        if r['status'] == "Success":
            lines.append(f"│  Trip:          {r['trip']} mi (DHO: {r['dho']} / DHD: {r['dhd']})")
            lines.append(f"│  Pricing:       Age: {r['age']} | Rate: {r['rate']} | RPM: {r['rpm']}")
            lines.append(f"│  Company:       {r['company']} (Phone: {r['phone']} | Email: {r['email']})")
            lines.append(f"│  Specs:         Weight: {r['weight']} | Length: {r['length']} | Factoring: {r['factoring']}")
            if r['ref_id']:
                lines.append(f"│  Ref ID:        {r['ref_id']}")
            
            # Formatted DAT Directory details block
            if r.get("dir_docket") or r.get("dir_dot_number"):
                lines.append(f"│  DAT Directory details:")
                lines.append(f"│    Docket: {r.get('dir_docket', '')} | DOT: {r.get('dir_dot_number', '')}")
                lines.append(f"│    Entity Type: {r.get('dir_gen_entity_type', '')} | Coverage To: {r.get('dir_active_ins_coverage_to', '')}")
                lines.append(f"│    Credit Score: {r.get('dir_credit_score', '')} | Insurance Carrier: {r.get('dir_active_ins_carrier', '')}")

            # Deep-extraction detail block (rate, route, truck, contact, office)
            deep = r.get("_deep") or {}
            if deep:
                def _put(label, path, width=20):
                    val = deep_value(r, path)
                    if val not in ("", None):
                        lines.append(f"│    {(label + ':').ljust(width)}{val}")

                rate = deep.get("rate") or {}
                if any(v for v in rate.values()):
                    lines.append("│  RATE (detail):")
                    _put("Total", "rate.rateTotal")
                    _put("Total trip", "rate.totalTrip")
                    _put("Rate/mile", "rate.ratePerMile")
                    _put("Total cost", "rate.totalCost")
                    _put("Spot rate", "rate.spotRate")
                    _put("Spot rate/mile", "rate.spotRatePerMile")
                    _put("Market", "rate.rateMarket")
                    _put("Range", "rate.rateRange")
                    _put("Contract rate", "rate.contractRate")

                route = deep.get("route") or {}
                if route.get("stops") or route.get("totalMiles"):
                    lines.append("│  ROUTE (detail):")
                    _put("Total miles", "route.totalMiles")
                    _put("Stops", "route.stopsFlat")

                truck = deep.get("truck") or {}
                if any(v for v in truck.values()):
                    lines.append("│  TRUCK (detail):")
                    _put("Truck", "truck.truckType")
                    _put("Equipment", "truck.equipmentType")
                    _put("Length", "truck.length")
                    _put("Weight", "truck.weight")

                contact = deep.get("contact") or {}
                if contact.get("email") or contact.get("phone") or contact.get("name"):
                    lines.append("│  CONTACT (detail):")
                    _put("Name", "contact.name")
                    _put("Phone", "contact.phone")
                    _put("Email", "contact.email")
                    _put("All emails", "contact.emailsFlat")

                office = deep.get("office") or {}
                if office.get("name") or office.get("addressRaw") or office.get("phone"):
                    lines.append("│  OFFICE:")
                    _put("Name", "office.name")
                    _put("Address", "office.addressRaw")
                    _put("City", "office.address.city")
                    _put("State", "office.address.state")
                    _put("ZIP", "office.address.zip")
                    _put("Phone", "office.phone")
                    _put("Fax", "office.fax")

                extra = deep.get("additionalData") or {}
                if extra:
                    lines.append(f"│  ADDITIONAL DATA ({len(extra)} unmapped fields):")
                    for k, v in list(extra.items())[:25]:
                        lines.append(f"│    {(k + ':').ljust(24)}{v}")

            ex = r.get("_extraction") or {}
            if ex:
                lines.append(f"│  EXTRACTION:   {ex.get('status', 'n/a')} · {ex.get('fieldsExtracted', 0)} fields")
                if ex.get("sectionsVisited"):
                    lines.append(f"│    Sections:   {', '.join(ex['sectionsVisited'])}")
                if ex.get("sectionsFailed"):
                    lines.append(f"│    Failed:     {', '.join(ex['sectionsFailed'])}")
                if ex.get("errors"):
                    err_txt = " | ".join(
                        "{}: {}".format(e.get("stage", ""), e.get("message", ""))
                        for e in ex["errors"]
                    )
                    lines.append(f"│    Errors:     {err_txt}")
        elif r['error']:
            lines.append(f"│  Error:         {r['error']}")
        lines.append(f"│  Timestamp:     {r['timestamp']}")
        lines.append("└" + "─"*78)
        lines.append("")
        
    return "\n".join(lines)

def get_export_xlsx(rows, export_type):
    from openpyxl.cell import WriteOnlyCell
    
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title="DAT Loads")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Pre-define fills for caching skipped/failed rows
    skipped_fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    failed_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    
    # Header and source key are bound together so the two lists can never drift.
    # They previously did: five headers had no matching value, which silently
    # shifted every column after "factoring" under the wrong heading.
    legacy_columns = [
        ("timestamp", "timestamp"),
        ("age", "age"),
        ("rate", "rate"),
        ("ratePerMile", "ratePerMile"),
        ("tripMiles", "tripMiles"),
        ("origin", "origin"),
        ("destination", "destination"),
        ("deadheadOrigin", "deadheadOrigin"),
        ("deadheadDest", "deadheadDest"),
        ("equipmentType", "equipmentType"),
        ("weight", "weight"),
        ("length", "length"),
        ("loadType", "loadType"),
        ("company", "company"),
        ("truckType", "truckType"),
        ("factoring", "factoring"),
        ("pickupDate", "pickupDate"),
        ("pickupTime", "pickupTime"),
        ("deliveryTime", "deliveryTime"),
        ("phone", "phone"),
        ("email", "email"),
        ("isEmpty", "isEmpty"),
        ("rowStatus", "rowStatus"),
        ("_captureKey", "_captureKey"),
        ("listIndex", "listIndex"),
        ("_apOrigin", "_apOrigin"),
        ("_apDest", "_apDest"),
        ("_apEq", "_apEq"),
        ("_apLt", "_apLt"),
        ("Docket", "dir_docket"),
        ("DotNumber", "dir_dot_number"),
        ("CreditScore", "dir_credit_score"),
    ]
    headers = [h for h, _ in legacy_columns]

    include_deep = rows_have_deep(rows)
    if include_deep:
        headers = headers + DEEP_EXPORT_HEADERS

    # Write headers
    header_cells = []
    for h in headers:
        cell = WriteOnlyCell(ws, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin
        header_cells.append(cell)
    ws.append(header_cells)
    
    col_widths = [len(h) for h in headers]
    align_left_cols = {15, 24, 27}
    
    for r in rows:
        row_data = [
            (r.get("isEmpty", False) if key == "isEmpty" else r.get(key, ""))
            for _, key in legacy_columns
        ]

        if include_deep:
            row_data.extend(deep_value(r, path) for _, path in DEEP_EXPORT_COLUMNS)

        status_val = str(r.get("rowStatus", ""))
        is_success = status_val == "Active" or status_val == "Success" or "success" in status_val.lower() or status_val == ""
        
        if not is_success:
            row_fill = failed_fill if ("failed" in status_val.lower() or "error" in status_val.lower()) else skipped_fill
            row_cells = []
            for col_idx, val in enumerate(row_data, 1):
                cell = WriteOnlyCell(ws, value=val)
                cell.border = border_thin
                cell.fill = row_fill
                cell.alignment = align_left if col_idx in align_left_cols else align_center
                row_cells.append(cell)
            ws.append(row_cells)
        else:
            ws.append(row_data)
            
        for col_idx, val in enumerate(row_data):
            val_str = str(val or '')
            if len(val_str) > col_widths[col_idx]:
                col_widths[col_idx] = len(val_str)
                
    for col_idx, width in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max(width + 3, 10), 40)
        
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    temp_file_path = temp_file.name
    temp_file.close()
    wb.save(temp_file_path)
    return temp_file_path

@app.route('/export', methods=['GET'])
def export_dataset():
    try_id = request.args.get("try_id")
    export_type = request.args.get("type", "all")  # all, success, skipped, failed
    export_format = request.args.get("format", "json")  # json, csv, txt, xlsx
    
    if not try_id:
        return jsonify({"success": False, "error": "try_id is required"}), 400
        
    rows = get_export_data(try_id, export_type)
    filename_base = f"dat-export-{export_type}-{try_id[:8]}-{int(time.time())}"
    
    if export_format == "json":
        res_data = json.dumps(rows, indent=2)
        response = make_response(res_data)
        response.headers['Content-Type'] = 'application/json'
        response.headers['Content-Disposition'] = f'attachment; filename={filename_base}.json'
        return response
        
    elif export_format == "csv":
        res_data = get_export_csv(rows)
        response = make_response(res_data)
        response.headers['Content-Type'] = 'text/csv'
        response.headers['Content-Disposition'] = f'attachment; filename={filename_base}.csv'
        return response
        
    elif export_format == "txt":
        res_data = get_export_txt(rows, export_type)
        response = make_response(res_data)
        response.headers['Content-Type'] = 'text/plain'
        response.headers['Content-Disposition'] = f'attachment; filename={filename_base}.txt'
        return response
        
    elif export_format == "xlsx":
        try:
            temp_xlsx_path = get_export_xlsx(rows, export_type)
            with open(temp_xlsx_path, "rb") as f:
                file_data = f.read()
            try:
                os.unlink(temp_xlsx_path)
            except Exception:
                pass
            response = make_response(file_data)
            response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            response.headers['Content-Disposition'] = f'attachment; filename={filename_base}.xlsx'
            return response
        except Exception as e:
            log_event("EXPORT_ERROR", {"try_id": try_id, "format": export_format, "error": str(e)})
            return jsonify({"success": False, "error": str(e)}), 500
            
    return jsonify({"success": False, "error": "Invalid format"}), 400

@app.route('/merge', methods=['POST'])
def merge():
    wal = load_json(WAL_FILE, {})
    try_id = wal.get("try_id")
    if not try_id:
        return jsonify({"success": False, "error": "No active try found to merge"})

    try_dir = os.path.join(BASE_DIR, f"try_{try_id}")
    if not os.path.exists(try_dir):
        return jsonify({"success": False, "error": f"Directory {try_dir} not found"})

    # Load WAL from the try directory to ensure up‑to‑date batch info
    wal_path = os.path.join(try_dir, "wal.json")
    wal = load_json(wal_path, {})

    all_results = []
    for f in sorted(os.listdir(try_dir)):
        if f.endswith(".json") and not f.startswith("skipped_") and f != "dataset.json" and f != "dataset_stats.json":
            batch_data = load_json(os.path.join(try_dir, f), [])
            all_results.extend(batch_data)

    merged_file = os.path.join(BASE_DIR, f"merged_try_{try_id}.json")
    save_json(merged_file, all_results)
    log_event("MERGE", {"try_id": try_id, "total_records": len(all_results), "merged_file": merged_file})

    return jsonify({"success": True, "merged_file": merged_file, "total_records": len(all_results)})

if __name__ == '__main__':
    print(f"Starting server in {BASE_DIR}")
    print("Ensure to run this when using the DAT Load Extractor V2.")
    app.run(port=5000, debug=True)
